"""Convert functional testcase Markdown to MeterSphere import Excel.

Usage:
    python md_to_metersphere_excel.py --feature-slug product-asset-tag
    python md_to_metersphere_excel.py --feature-slug assets-card-customization --sprint OBIS-20260622-20260703
    python md_to_metersphere_excel.py --input path/to.md --output out.xlsx
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[4]
if str(REPO_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "scripts"))

from feature_paths import testcases_md_path, testcases_xlsx_path  # noqa: E402

DEFAULT_TEMPLATE = REPO_ROOT / "docs/templates/import_ms_excel_case.xlsx"

HEADERS = [
    "用例名称",
    "所属模块",
    "标签",
    "前置条件",
    "步骤描述",
    "预期结果",
    "编辑模式",
    "备注",
    "用例类型",
    "用例等级",
    "用例分组",
    "用例描述",
    "Case Review",
    "Is Regression",
]

MD_COLUMNS = [
    "用例名称",
    "所属模块",
    "标签",
    "用例等级",
    "前置条件",
    "步骤描述",
    "预期结果",
    "备注",
]


def normalize_cell(text: str) -> str:
    text = text.replace("<br>", "\n")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return text.strip()


def parse_md_table(path: Path) -> list[dict[str, str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    header_idx = None
    for i, line in enumerate(lines):
        if line.strip().startswith("|") and "用例名称" in line and "所属模块" in line:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(f"No testcase table found in {path}")

    headers = [h.strip() for h in lines[header_idx].strip().strip("|").split("|")]
    if headers != MD_COLUMNS:
        raise ValueError(f"Unexpected MD columns: {headers}")

    rows: list[dict[str, str]] = []
    for line in lines[header_idx + 2 :]:
        if not line.strip().startswith("|"):
            break
        if re.match(r"^\|\s*-+\s*\|", line):
            continue
        cells = [normalize_cell(c) for c in line.strip().strip("|").split("|")]
        if len(cells) != len(headers):
            raise ValueError(f"Column count mismatch in row: {line[:80]}...")
        rows.append(dict(zip(headers, cells)))
    return rows


def write_excel(rows: list[dict[str, str]], template: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)

    wb = openpyxl.load_workbook(output)
    ws = wb.active

    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col, value=header)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_idx, case in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=case["用例名称"])
        ws.cell(row=row_idx, column=2, value=case["所属模块"])
        ws.cell(row=row_idx, column=3, value=case["标签"])
        ws.cell(row=row_idx, column=4, value=case["前置条件"])
        ws.cell(row=row_idx, column=5, value=case["步骤描述"])
        ws.cell(row=row_idx, column=6, value=case["预期结果"])
        ws.cell(row=row_idx, column=7, value="STEP")
        ws.cell(row=row_idx, column=8, value=case.get("备注", ""))
        ws.cell(row=row_idx, column=9, value="")
        ws.cell(row=row_idx, column=10, value=case["用例等级"])
        ws.cell(row=row_idx, column=11, value="")
        ws.cell(row=row_idx, column=12, value="")
        ws.cell(row=row_idx, column=13, value="")
        ws.cell(row=row_idx, column=14, value="")

    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert testcase MD to MeterSphere Excel")
    parser.add_argument("--feature-slug", help="Feature slug, e.g. product-asset-tag")
    parser.add_argument(
        "--sprint",
        help="Sprint id (optional if slug is registered), e.g. OBIS-20260622-20260703",
    )
    parser.add_argument("--input", type=Path, help="Input markdown path")
    parser.add_argument("--output", type=Path, help="Output xlsx path")
    parser.add_argument(
        "--template",
        type=Path,
        default=DEFAULT_TEMPLATE,
        help="MeterSphere import template xlsx",
    )
    args = parser.parse_args()

    if args.input:
        input_path = args.input
    elif args.feature_slug:
        try:
            input_path = testcases_md_path(args.feature_slug, args.sprint, REPO_ROOT)
        except KeyError as exc:
            parser.error(str(exc))
    else:
        parser.error("Provide --feature-slug or --input")

    if args.output:
        output_path = args.output
    elif args.feature_slug:
        output_path = testcases_xlsx_path(args.feature_slug, args.sprint, REPO_ROOT)
    else:
        output_path = input_path.with_suffix(".xlsx")

    rows = parse_md_table(input_path)
    write_excel(rows, args.template, output_path)
    print(f"Wrote {len(rows)} cases to {output_path}")


if __name__ == "__main__":
    main()
